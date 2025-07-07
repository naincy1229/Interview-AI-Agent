# file: app.py

import streamlit as st
import ollama
import time
import pyttsx3
import speech_recognition as sr
from datetime import datetime
from fpdf import FPDF
from docx import Document

st.set_page_config(page_title="Mock Interview AI", layout="centered")

# Initialize TTS engine
tts_engine = pyttsx3.init()
tts_engine.setProperty('rate', 170)

# Predefined question banks
QUESTION_BANK = {
    "HR": [
        "Tell me about yourself.",
        "What are your strengths and weaknesses?",
        "Why should we hire you?",
    ],
    "Technical": [
        "Explain a challenging coding problem you solved.",
        "How do you handle version control in teams?",
        "What is the time complexity of quicksort?",
    ],
    "Behavioral": [
        "Describe a time you faced conflict in a team.",
        "How do you prioritize tasks under pressure?",
        "Give an example of when you showed leadership.",
    ]
}

if 'category' not in st.session_state:
    st.session_state.category = "HR"
if 'questions' not in st.session_state:
    st.session_state.questions = QUESTION_BANK[st.session_state.category]
if 'q_index' not in st.session_state:
    st.session_state.q_index = 0
if 'logs' not in st.session_state:
    st.session_state.logs = []
if 'resume_text' not in st.session_state:
    st.session_state.resume_text = ""

st.title("🧠 AI Mock Interview Coach")

# Upload resume
resume_file = st.file_uploader("📄 Upload Your Resume (DOCX only)", type=['docx'])
if resume_file:
    doc = Document(resume_file)
    resume_text = "\n".join([para.text for para in doc.paragraphs])
    st.session_state.resume_text = resume_text
    st.success("✅ Resume loaded. Interview will be personalized.")

# Question Category
st.selectbox("📂 Choose Question Category:", options=list(QUESTION_BANK.keys()), key='category', on_change=lambda: (st.session_state.update({'questions': QUESTION_BANK[st.session_state.category], 'q_index': 0, 'logs': []})))

question = st.session_state.questions[st.session_state.q_index]
st.markdown(f"### 💬 {question}")

# Timer
with st.expander("⏱️ Timer (3 mins)"):
    if st.button("Start Timer"):
        for sec in range(180, 0, -1):
            st.write(f"⏳ Time left: {sec}s")
            time.sleep(1)
            st.experimental_rerun()

# Record Audio
def record_audio():
    recognizer = sr.Recognizer()
    with sr.Microphone() as source:
        st.info("🎙 Speak now...")
        audio_data = recognizer.listen(source, phrase_time_limit=10)
    try:
        text = recognizer.recognize_google(audio_data)
        return text
    except sr.UnknownValueError:
        return "[Unrecognized speech]"
    except sr.RequestError:
        return "[Speech recognition service failed]"

user_answer = st.text_area("🧑 Your Answer:", height=150)

if st.button("🎙 Use Voice Input"):
    result = record_audio()
    st.write(f"🗣 You said: {result}")
    user_answer = result

if st.button("🎯 Get Feedback"):
    if user_answer.strip():
        with st.spinner("🤖 Thinking..."):
            prompt = f"""
You are a senior interviewer.
Use the resume below to personalize your feedback.

Resume:
{st.session_state.resume_text}

Question: {question}
Answer: {user_answer}

Provide:
1. Strengths
2. Areas to Improve
3. Rate the answer out of 10
4. A follow-up question
"""
            response = ollama.chat(
                model="llama2",
                messages=[{"role": "user", "content": prompt}]
            )
            feedback = response['message']['content']

        st.markdown("### 📝 Feedback:")
        st.write(feedback)

        # Voice output
        if st.button("🔊 Read Feedback Aloud"):
            tts_engine.say(feedback)
            tts_engine.runAndWait()

        # Save log
        log_entry = {
            "question": question,
            "answer": user_answer,
            "feedback": feedback,
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        }
        st.session_state.logs.append(log_entry)

        # Save to file
        filename = f"interview_log_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"
        with open(filename, "a", encoding="utf-8") as f:
            f.write(f"Q: {question}\nA: {user_answer}\n{feedback}\n{'-'*40}\n")

        if st.session_state.q_index < len(st.session_state.questions) - 1:
            st.session_state.q_index += 1
            st.experimental_rerun()
        else:
            st.success("✅ Interview completed!")
    else:
        st.warning("Please enter or record an answer first.")

# Follow-up Chat
st.markdown("---")
st.markdown("### 💬 Ask a follow-up about the feedback")
followup = st.text_input("Type your follow-up question")
if st.button("Ask AI"):
    if followup:
        with st.spinner("🤖 Responding..."):
            f_prompt = f"The candidate asked a follow-up based on previous feedback: {followup}. Answer as a helpful interviewer."
            f_response = ollama.chat(
                model="llama2",
                messages=[{"role": "user", "content": f_prompt}]
            )
            st.write(f_response['message']['content'])

# PDF Export
if st.button("📄 Download PDF Report"):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="AI Interview Feedback Report", ln=True, align="C")
    pdf.ln(10)
    for log in st.session_state.logs:
        pdf.multi_cell(0, 10, txt=f"Q: {log['question']}\nA: {log['answer']}\n{log['feedback']}\n{'-'*40}")
        pdf.ln(4)
    report_name = f"Interview_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    pdf.output(report_name)
    st.success(f"📄 PDF saved as {report_name}")
